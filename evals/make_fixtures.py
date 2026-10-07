"""Generate the test documents under evals/fixtures/.

DOCX files are built with python-docx (plus raw XML where python-docx has no API).
PDF and DOC copies of them are rendered with LibreOffice. `layout.pdf` is drawn directly
with PyMuPDF so its coordinates are exact. Outputs are committed; rerun only to change them.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "fixtures"
sys.path.insert(0, str(ROOT.parent / "src"))

from ai_doc_editor.soffice import convert, find_soffice  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = parse_xml(
        f'<w:hyperlink xmlns:w="{W_NS}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}"><w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
        f"<w:t>{text}</w:t></w:r></w:hyperlink>"
    )
    paragraph._p.append(link)


def _add_page_field(paragraph) -> None:
    paragraph._p.append(
        parse_xml(
            f'<w:fldSimple xmlns:w="{W_NS}" w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
        )
    )


def _add_text_box(paragraph, text: str) -> None:
    """Inline DrawingML text box (wps:wsp) holding one paragraph."""
    xml = f"""
<w:r xmlns:w="{W_NS}"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:drawing>
    <wp:inline distT="0" distB="0" distL="0" distR="0">
      <wp:extent cx="3600000" cy="540000"/>
      <wp:docPr id="100" name="TextBox 1"/>
      <a:graphic>
        <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
          <wps:wsp>
            <wps:cNvSpPr txBox="1"/>
            <wps:spPr>
              <a:xfrm><a:off x="0" y="0"/><a:ext cx="3600000" cy="540000"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
              <a:ln w="6350"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>
            </wps:spPr>
            <wps:txbx>
              <w:txbxContent>
                <w:p><w:r><w:rPr><w:i/></w:rPr><w:t>{text}</w:t></w:r></w:p>
              </w:txbxContent>
            </wps:txbx>
            <wps:bodyPr/>
          </wps:wsp>
        </a:graphicData>
      </a:graphic>
    </wp:inline>
  </w:drawing>
</w:r>"""
    paragraph._p.append(parse_xml(xml))


def _set_columns(section, num: int) -> None:
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = parse_xml(f'<w:cols xmlns:w="{W_NS}"/>')
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "708")


def make_report_en(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2)
    sec.header.paragraphs[0].text = "Acme Corp - Quarterly Report"
    footer_p = sec.footer.paragraphs[0]
    footer_p.add_run("Page ")
    _add_page_field(footer_p)

    doc.add_heading("Quarterly Report for Acme Corp", level=1)
    p = doc.add_paragraph("This report was prepared by ")
    p.add_run("Acme Corp").bold = True
    p.add_run(" for its ")
    p.add_run("board of directors").italic = True
    r = p.add_run(". Revenue grew strongly")
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(" in the third quarter. Details are on ")
    _add_hyperlink(p, "https://example.com/acme", "the investor site")
    p.add_run(".")

    doc.add_paragraph(
        "We recieve many questions about the merger, and we beleive the answers below "
        "adress the most common ones."
    )
    doc.add_heading("Highlights", level=2)
    for text in [
        "Revenue increased by 12% year over year.",
        "Acme Corp opened two new offices.",
        "Customer churn fell to 3.1%.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    for text in ["Approve the budget.", "Elect the new auditor.", "Close the meeting."]:
        doc.add_paragraph(text, style="List Number")

    table = doc.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    header = table.rows[0].cells
    merged = header[0].merge(header[1])
    merged.text = "Region"
    header[2].text = "Q2"
    header[3].text = "Q3"
    rows = [("North", "Acme North", "1.2", "1.4"), ("South", "Acme South", "0.8", "0.9")]
    for i, (a, b, c, d) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text = a, b, c
        cells[3].paragraphs[0].add_run(d).bold = True

    box_p = doc.add_paragraph()
    _add_text_box(box_p, "Note: figures are unaudited and provided by Acme Corp.")

    new_sec = doc.add_section()
    new_sec.orientation = WD_ORIENT.LANDSCAPE
    new_sec.page_width, new_sec.page_height = Cm(29.7), Cm(21)
    _set_columns(new_sec, 2)
    doc.add_heading("Appendix", level=2)
    for i in range(6):
        para = doc.add_paragraph(
            f"Appendix paragraph {i + 1}. Acme Corp keeps records of every transaction "
            "for at least seven years, as required by law."
        )
        para.paragraph_format.space_after = Pt(6)
    doc.save(path)


def make_contract_ru(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header.paragraphs[0].text = "Договор поставки № 17/2026"

    doc.add_heading("Договор поставки", level=1)
    p = doc.add_paragraph("ООО «Ромашка», в лице генерального директора ")
    p.add_run("Иванова Петра Сергеевича").bold = True
    p.add_run(", действующего на основании Устава, именуемое в дальнейшем «Поставщик», и ")
    p.add_run("ООО «Василёк»").bold = True
    p.add_run(", именуемое в дальнейшем «Покупатель», заключили настоящий договор.")
    doc.add_paragraph(
        "Поставщик обязуется передать товар в сроки, указаные в Приложении 1, "
        "а Покупатель обязуется принять и оплатить товар."
    )
    doc.add_paragraph(
        "Контактное лицо Поставщика: Иванов П. С., тел. +7 495 123-45-67, "
        "e-mail: ivanov@romashka.example."
    )
    doc.add_paragraph(
        "Контактное лицо Покупателя: Смирнова А. В., тел. +7 812 765-43-21, "
        "e-mail: smirnova@vasilek.example."
    )
    doc.add_heading("Реквизиты сторон", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Поставщик: ООО «Ромашка», ИНН 7701234567"
    table.cell(0, 1).text = "Покупатель: ООО «Василёк», ИНН 7809876543"
    table.cell(1, 0).text = "Подпись: ____________ Иванов П. С."
    table.cell(1, 1).text = "Подпись: ____________ Смирнова А. В."
    doc.save(path)


def make_layout_pdf(path: Path) -> None:
    """A4 page with text blocks at exact positions, an image and vector graphics."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    noto = pymupdf.Font("notos")
    page.insert_font(fontname="notos", fontbuffer=noto.buffer)
    noto_b = pymupdf.Font("notosbo")
    page.insert_font(fontname="notosbo", fontbuffer=noto_b.buffer)

    page.insert_textbox(
        pymupdf.Rect(72, 60, 523, 90),
        "Acme Corp Service Agreement",
        fontname="notosbo",
        fontsize=18,
        align=pymupdf.TEXT_ALIGN_CENTER,
    )
    page.draw_line((72, 96), (523, 96), color=(0.2, 0.2, 0.6), width=1.5)
    page.insert_textbox(
        pymupdf.Rect(72, 110, 523, 170),
        "This agreement is made between Acme Corp and the Client. Acme Corp will provide "
        "support services during business hours.",
        fontname="notos",
        fontsize=11,
    )
    page.insert_textbox(
        pymupdf.Rect(72, 180, 290, 300),
        "Left column. Payments are due within thirty days of the invoice date.",
        fontname="notos",
        fontsize=10,
    )
    page.insert_textbox(
        pymupdf.Rect(305, 180, 523, 300),
        "Правая колонка. Оплата производится в течение тридцати дней.",
        fontname="notos",
        fontsize=10,
    )
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
    for x in range(64):
        for y in range(64):
            pix.set_pixel(x, y, (x * 4, y * 4, 128))
    page.insert_image(pymupdf.Rect(72, 320, 172, 420), pixmap=pix)
    page.draw_rect(pymupdf.Rect(190, 320, 523, 420), color=(0, 0, 0), fill=(0.93, 0.93, 0.97))
    page.insert_textbox(
        pymupdf.Rect(200, 330, 513, 410),
        "Contact: support@acme.example, phone +1 555 0100.",
        fontname="notos",
        fontsize=10,
    )
    page.insert_textbox(
        pymupdf.Rect(72, 780, 523, 800),
        "Page 1",
        fontname="notos",
        fontsize=8,
        align=pymupdf.TEXT_ALIGN_RIGHT,
    )
    doc.save(path, garbage=4, deflate=True)


def make_sections_pdf(path: Path) -> None:
    """Two sections, every line its own text object (as browsers print), with room to reflow."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="notos", fontbuffer=pymupdf.Font("notos").buffer)
    page.insert_font(fontname="notosbo", fontbuffer=pymupdf.Font("notosbo").buffer)
    y = 80.0

    def heading(text: str) -> None:
        nonlocal y
        page.insert_text((60, y), text, fontname="notosbo", fontsize=16)
        y += 30

    def paragraph(lines: list[str]) -> None:
        nonlocal y
        for line in lines:
            page.insert_text((60, y), line, fontname="notos", fontsize=10.5)
            y += 16
        y += 10

    heading("1. About This Certification")
    paragraph(
        [
            "The certification validates that an individual can design, build and deliver",
            "production-grade AI solutions. It is intended for practitioners working in an",
            "architect role who select models and integrate them into enterprise systems.",
        ]
    )
    paragraph(
        [
            "This guide is the authoritative reference for candidates preparing to sit the",
            "exam. Read it in full before scheduling your exam.",
        ]
    )
    y += 40
    heading("2. Purpose of the Credential")
    paragraph(
        [
            "The credential provides an independent assessment of the skills required to",
            "architect solutions in production environments.",
        ]
    )
    doc.save(path, garbage=4, deflate=True)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_report_en(OUT / "report_en.docx")
    make_contract_ru(OUT / "contract_ru.docx")
    make_layout_pdf(OUT / "layout.pdf")
    make_sections_pdf(OUT / "sections.pdf")
    if find_soffice() is None:
        print("soffice not found: skipping PDF and DOC renders of the DOCX fixtures")
        return
    for name in ("report_en", "contract_ru"):
        convert(OUT / f"{name}.docx", "pdf", OUT)
    convert(OUT / "contract_ru.docx", "doc", OUT)
    print("fixtures written to", OUT)


if __name__ == "__main__":
    main()
